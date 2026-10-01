#!/usr/bin/env python3
"""
Mechanical extractor for SAP Readiness Check for SAP S/4HANA exports.

SAP generates the report .docx with a fixed, numbered heading structure
(chapters 1-10, stable across every real customer export checked: identical
wording and numbering whether the source system is a small test box or a
large production landscape). This script exploits that stability: it walks
the document body in order, splits it into chapters/subchapters by heading
style, and captures each chapter's paragraph text and tables as structured
data. It does not interpret anything — severity judgment, "is this a real
risk" calls, and cross-references into a Solution Architecture Document all
belong to the model, guided by knowledge/methodology/brownfield-conversion/readiness-check-interpretation.md (see the readiness-analysis Skill).

A real client export is never just the .docx — the Readiness Check
dashboard bundles it with ~20-25 supporting .xlsx exports (Relevant
Simplification Items, Technical Properties, Add-On Compatibilities, HANA
sizing tables, Interface Impact Analysis per protocol, Financial Data
tables, etc.). Those .xlsx files carry the real underlying numbers as
clean structured columns (Category, Effort Ranking, the "(default)" vs
"(calculated)" distinction, per-item Errors/Warnings counts) — reading
them beats re-parsing the same numbers out of docx table cell text, which
is lossier and error-prone. This script ingests both.

Usage:
    # Single .docx only (supporting_data will be empty):
    python3 extract_readiness_check.py <report.docx> <output.json>

    # A folder (or the unzipped contents of the client's export .zip)
    # containing the .docx plus its supporting .xlsx exports — this is the
    # preferred mode for a real client export:
    python3 extract_readiness_check.py <export_folder/> <output.json>

Output JSON shape:
{
  "chapters": [
    {
      "number": "3.1.1",
      "level": 3,
      "heading": "Relevant Simplification Items",
      "paragraphs": ["...", "..."],
      "tables": [
        {"headers": ["LoB", "Title", ...], "rows": [["...", "..."], ...]}
      ],
      "gap_flags": ["No custom code analysis results have been uploaded."],
      "clean_result_flags": [],
      "structurally_empty": false
    },
    ...
  ],
  "supporting_data": {
    "RelevantSimplificationItems": {
      "<sheet name>": {"headers": [...], "row_count": 108, "rows": [[...], ...]}
    },
    "HANATopDatabaseTablesByDiskUsage": {
      "<sheet name>": {"headers": [...], "row_count": 1961, "sidecar": "rc_data_supporting/HANATopDatabaseTablesByDiskUsage__SAP_Document_Export.json"}
    },
    ...
  }
}

gap_flags: paragraphs matching known "no data collected / check could not be
performed" phrasing are surfaced separately, because in real reports entire
chapters (Custom Code, Data Volume Management, Flat File interfaces, Add-On
Compatibility) routinely come back empty for this reason — that is itself a
planning gap to report, not a clean bill of health, and it's easy to miss
buried in an otherwise-empty chapter.

clean_result_flags: paragraphs matching known "the check ran and found
nothing" phrasing (e.g. "No Inconsistencies were Found" in Financial Data
Quality) — kept in a separate field from gap_flags on purpose, since it is
the opposite finding (clean, not missing) despite often being the sole,
short content of the chapter.

structurally_empty: true when a chapter has no tables and no substantive
paragraph text, regardless of whether any known gap phrase matched. The
phrase list can't cover every wording SAP uses across report versions (a
real export was found with an empty Flat File Interfaces section that
matched no known phrase at all) — this is a second, content-shape-based
signal to catch what the phrase list misses. It is a prompt to check the
section by hand, not itself proof of a data-collection gap: some sections
(e.g. a short "see linked report" chapter) are legitimately this short.

supporting_data keys are the source filename without its extension, so
they line up with the client export's own naming (e.g. a file named
"RelevantSimplificationItems.xlsx" produces the key
"RelevantSimplificationItems"). Every worksheet in each file is captured;
which sheet/column names to expect is interpretation guidance, not
something this script hardcodes, since SAP has changed column layouts

A sheet with more than SIDECAR_ROW_THRESHOLD rows is written to its own
JSON file under "<output>_supporting/" instead of being inlined — the
combined output otherwise balloons well past what's practical to read in
one pass (a real 24-file export produced a 4.7MB single JSON: several
supporting sheets alone run 1,600-3,400 rows — HANA sizing tables,
interface impact analyses, Fiori app recommendations). Inlined sheets
carry a "rows" key as before; sidecar sheets instead carry a "sidecar"
key (a path relative to the output JSON's directory) and no "rows" — read
the sidecar file only when that specific chapter's interpretation
actually needs it (e.g. Chapter 6 needs the HANA sizing sidecars; most
other chapters don't), not preemptively. Every sheet, inlined or not,
still gets a "row_count" so its scale is visible without opening it.
across Readiness Check tool versions.
"""
import sys
import os
import glob
import json
import re
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl

HEADING_NUM_RE = re.compile(r"^(\d+(?:\.\d+)*)\s+(.*)$")

GAP_PHRASES = [
    "could not be performed",
    "could not find any",
    "no data was collected",
    "no data collected",
    "no custom code analysis",
    "not been uploaded",
    "there is no document type",
    "no sap innovative business solutions available",
    "please follow the steps in sap note",
    "no simplification item with relevance to be checked",
]

# Phrases that report a genuinely clean result (the check ran and found
# nothing), not a data-collection gap. Kept out of GAP_PHRASES/gap_flags on
# purpose: a real export was found where this exact phrase was the sole
# content of two Financial Data Quality "Inconsistencies" subchapters, and
# lumping it in with gap_flags would report a clean bill of health as an
# unresolved gap ("impact unknown") — the opposite of what the report says.
CLEAN_RESULT_PHRASES = [
    "no inconsistencies were found",
]

# A chapter this short (in characters, across its substantive paragraph
# text) with no tables is treated as structurally empty even without a
# phrase match.
STRUCTURALLY_EMPTY_CHAR_THRESHOLD = 40

# Paragraphs matching these are boilerplate/navigation, not real findings —
# excluded before measuring how much substantive text a chapter has. A real
# export was found with a Flat File Interfaces section containing only a
# "see the linked report" sentence plus a bare URL and no table at all; that
# reads as non-empty by a naive length check but is exactly the kind of
# never-actually-analyzed section the structural-emptiness signal exists to
# catch.
BOILERPLATE_PARAGRAPH_RE = re.compile(
    r"^https?://|^link to the (complete|full) list", re.IGNORECASE
)

# Sheets with more rows than this are written to a sidecar file instead of
# being inlined in the main output JSON (see module docstring).
SIDECAR_ROW_THRESHOLD = 150


def iter_block_items(doc):
    """Yield paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def table_to_dict(table):
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return {"headers": [], "rows": []}
    return {"headers": rows[0], "rows": rows[1:]}


def extract_chapters(docx_path):
    doc = docx.Document(docx_path)
    chapters = []
    current = None

    # First pass: find the index of the last occurrence of each numbered heading text among Heading-styled paragraphs,
    # so we only start a new chapter at the real section, not an earlier incidental mention.
    heading_last_index = {}
    for idx, item in enumerate(iter_block_items(doc)):
        if isinstance(item, Paragraph) and item.style.name.startswith("Heading"):
            text = item.text.strip()
            m = HEADING_NUM_RE.match(text)
            if m:
                heading_last_index[text] = idx

    for idx, item in enumerate(iter_block_items(doc)):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if item.style.name.startswith("Heading"):
                m = HEADING_NUM_RE.match(text)
                if m and heading_last_index.get(text) == idx:
                    number, heading = m.groups()
                    level = number.count(".") + 1
                    current = {
                        "number": number,
                        "level": level,
                        "heading": heading,
                        "paragraphs": [],
                        "tables": [],
                        "gap_flags": [],
                        "clean_result_flags": [],
                    }
                    chapters.append(current)
                    continue
            if current is not None and text:
                current["paragraphs"].append(text)
                lower = text.lower()
                if any(phrase in lower for phrase in CLEAN_RESULT_PHRASES):
                    current["clean_result_flags"].append(text)
                else:
                    for phrase in GAP_PHRASES:
                        if phrase in lower:
                            current["gap_flags"].append(text)
                            break
        elif isinstance(item, Table):
            if current is not None:
                current["tables"].append(table_to_dict(item))

    for chapter in chapters:
        substantive = [p for p in chapter["paragraphs"] if not BOILERPLATE_PARAGRAPH_RE.match(p.strip())]
        char_count = sum(len(p) for p in substantive)
        chapter["structurally_empty"] = (
            not chapter["tables"]
            and char_count < STRUCTURALLY_EMPTY_CHAR_THRESHOLD
            and not chapter["clean_result_flags"]
        )

    return chapters


def extract_xlsx(xlsx_path):
    """Read every sheet of one supporting .xlsx export into {sheet_name: {headers, rows}}.

    Deliberately NOT read_only=True: several real SAP Readiness Check .xlsx
    exports carry a stale <dimension> tag in the sheet XML that under-declares
    the actual row count (e.g. claims 1 row when 109 are present) — openpyxl's
    read_only mode trusts that tag and silently truncates iteration to it,
    while normal mode scans the real rows regardless of what the tag claims.
    These files are small enough (tens of KB to a few hundred KB) that the
    extra memory of loading them normally is a non-issue.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            sheets[sheet_name] = {"headers": [], "rows": []}
            continue
        headers = ["" if v is None else str(v) for v in header_row]
        rows = []
        for row in rows_iter:
            if row is None or all(v is None for v in row):
                continue
            rows.append(["" if v is None else v for v in row])
        sheets[sheet_name] = {"headers": headers, "rows": rows}
    return sheets


def resolve_inputs(input_path):
    """Return (docx_path, list_of_xlsx_paths) whether input_path is a .docx file or a folder."""
    if os.path.isdir(input_path):
        docx_candidates = glob.glob(os.path.join(input_path, "*.docx"))
        docx_candidates = [p for p in docx_candidates if not os.path.basename(p).startswith("~$")]
        if not docx_candidates:
            print(f"No .docx file found in {input_path}", file=sys.stderr)
            sys.exit(1)
        if len(docx_candidates) > 1:
            print(f"Multiple .docx files found in {input_path}, using the first: {docx_candidates[0]}", file=sys.stderr)
        docx_path = docx_candidates[0]
        xlsx_paths = sorted(glob.glob(os.path.join(input_path, "*.xlsx")))
        xlsx_paths = [p for p in xlsx_paths if not os.path.basename(p).startswith("~$")]
        return docx_path, xlsx_paths
    else:
        return input_path, []


def sanitize_for_filename(name):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_") or "sheet"


def split_out_sidecars(supporting_data, out_path):
    """Move any sheet over SIDECAR_ROW_THRESHOLD rows into its own JSON file.

    Returns the (possibly-rewritten) supporting_data and the list of sidecar
    file paths written, for reporting.
    """
    out_dir = os.path.dirname(os.path.abspath(out_path))
    out_stem = os.path.splitext(os.path.basename(out_path))[0]
    sidecar_dir = os.path.join(out_dir, f"{out_stem}_supporting")
    sidecar_paths = []

    for file_key, sheets in supporting_data.items():
        for sheet_name, table in sheets.items():
            row_count = len(table["rows"])
            table["row_count"] = row_count
            if row_count > SIDECAR_ROW_THRESHOLD:
                os.makedirs(sidecar_dir, exist_ok=True)
                sidecar_filename = f"{sanitize_for_filename(file_key)}__{sanitize_for_filename(sheet_name)}.json"
                sidecar_path = os.path.join(sidecar_dir, sidecar_filename)
                with open(sidecar_path, "w") as sf:
                    json.dump({"headers": table["headers"], "rows": table["rows"]}, sf, ensure_ascii=False)
                table["sidecar"] = os.path.relpath(sidecar_path, out_dir)
                del table["rows"]
                sidecar_paths.append(sidecar_path)

    return supporting_data, sidecar_paths


def extract(input_path, out_path):
    docx_path, xlsx_paths = resolve_inputs(input_path)

    chapters = extract_chapters(docx_path)

    supporting_data = {}
    for xlsx_path in xlsx_paths:
        key = os.path.splitext(os.path.basename(xlsx_path))[0]
        try:
            supporting_data[key] = extract_xlsx(xlsx_path)
        except Exception as e:
            print(f"Warning: failed to read {xlsx_path}: {e}", file=sys.stderr)

    supporting_data, sidecar_paths = split_out_sidecars(supporting_data, out_path)

    with open(out_path, "w") as f:
        json.dump({"chapters": chapters, "supporting_data": supporting_data}, f, indent=2, ensure_ascii=False)

    print(f"Extracted {len(chapters)} chapter/subchapter sections to {out_path}")
    gap_count = sum(1 for c in chapters if c["gap_flags"])
    print(f"{gap_count} sections contain a data-collection-gap phrase — check these explicitly, don't assume they're clean.")
    structural_count = sum(1 for c in chapters if c["structurally_empty"] and not c["gap_flags"])
    if structural_count:
        print(f"{structural_count} additional section(s) are structurally empty (no table, no real paragraph text) "
              f"but matched no known gap phrase — check these too, the phrase list doesn't cover every SAP wording.")
    if xlsx_paths:
        print(f"Ingested {len(supporting_data)} supporting .xlsx export(s): {', '.join(supporting_data.keys())}")
        if sidecar_paths:
            print(f"{len(sidecar_paths)} sheet(s) with more than {SIDECAR_ROW_THRESHOLD} rows were written to "
                  f"sidecar files instead of being inlined (see the 'sidecar' key in supporting_data) — "
                  f"read only the ones the chapter you're interpreting actually needs.")
    else:
        print("No supporting .xlsx exports found/provided — pass the export folder (not just the .docx) "
              "when the client's export includes them, so Simplification Item Category/Effort Ranking/"
              "Errors-Warnings and other findings come from clean structured columns instead of docx table text.")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    extract(sys.argv[1], sys.argv[2])


if __name__ == "__main__":
    main()
