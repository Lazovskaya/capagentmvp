#!/usr/bin/env python3
"""
Generic tabular extractor for a client's Application Map / interface
inventory. Every client formats this differently (columns, sheet names,
taxonomy for "which business area this app belongs to") — this script
does not assume any particular structure. It just turns each sheet into a
list of header->value dicts, using the first non-empty row as the header
row. The semantic work (which application supports which business
process, using whatever "Business Area"/"Business Domain"/"Functionalities"
column the client's file happens to have) is the model's job, guided by
references/scoping_methodology.md §6 — not this script's.

Usage:
    python3 extract_application_map.py <application_map.xlsx> <output.json> [--sheets "Sheet1,Sheet2"]

--sheets: comma-separated sheet names to extract. Omit to extract every
sheet in the workbook.
"""
import sys
import json
import argparse
import openpyxl


def first_header_row(rows):
    for row in rows:
        if row and any(c is not None and str(c).strip() for c in row):
            return row
    return None


def extract_sheet(ws):
    rows_iter = ws.iter_rows(values_only=True)
    rows = list(rows_iter)
    header = first_header_row(rows)
    if header is None:
        return []
    header_idx = rows.index(header)
    clean_header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header)]
    out = []
    for row in rows[header_idx + 1:]:
        if row is None or all(c is None for c in row):
            continue
        record = {}
        for h, v in zip(clean_header, row):
            record[h] = v
        out.append(record)
    return out


def extract(xlsx_path, out_path, sheet_names=None):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    targets = sheet_names if sheet_names else wb.sheetnames

    result = {}
    for name in targets:
        if name not in wb.sheetnames:
            stripped_match = next((s for s in wb.sheetnames if s.strip() == name.strip()), None)
            if stripped_match:
                name = stripped_match
        if name not in wb.sheetnames:
            print(f"Warning: sheet '{name}' not found, skipping. Available: {wb.sheetnames}")
            continue
        result[name] = extract_sheet(wb[name])

    with open(out_path, "w") as f:
        json.dump(result, f, indent=2, ensure_ascii=False, default=str)

    for name, rows in result.items():
        print(f"Sheet '{name}': {len(rows)} rows extracted")
    print(f"Wrote {out_path}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("out_path")
    parser.add_argument("--sheets", default=None, help="Comma-separated sheet names; omit for all sheets")
    args = parser.parse_args()
    sheet_names = args.sheets.split(",") if args.sheets else None
    extract(args.xlsx_path, args.out_path, sheet_names)


if __name__ == "__main__":
    main()
