#!/usr/bin/env python3
"""
Mechanical builder for the Business Process Scope deliverable as a real
.xlsx workbook, matching the shape of a real EPAM deliverable (ADNOC L&S
ERP Transformation, Section 2.2 Business Process Scope): one "Summary"
sheet plus one sheet per Business Area/LoB.

This script does not decide Changed/Unchanged/complexity — that judgment
is the model's job (see references/scoping_methodology.md). It only takes
the already-decided structured data and lays it out as a workbook,
consistently, every time — so the deliverable is always an Excel file,
not just prose or markdown.

Usage:
    python3 build_scope_workbook.py <input.json> <output.xlsx>

Input JSON shape:
{
  "title": "Client Name — Business Process Scope",
  "areas": {
    "Finance": [
      {"stream": "Accounting and Financial Close", "process": "Asset Accounting",
       "status": "Changed", "complexity": "High",
       "driving_item": "S4TWL - ASSET ACCOUNTING", "note": "Potentially High effort ranking"},
      ...
    ],
    "Manufacturing": [ ... ]
  },
  "notes": ["Planning-level draft...", "Analysis ran on Test System..."]
}

Rows may omit "driving_item"/"note" (written as blank cells). "status"
should be one of Changed / Unchanged / To Assess / Out of Scope — any
other value is written as-is but won't get the status column's color
coding.

Optional "applications" field per row (a string, or a list of strings —
e.g. ["SAP ECC EHP8", "READSOFT KOFAX PD-AP"]) adds a "Supporting
Application(s)" column, populated from cross-referencing the client's own
Application Map (see extract_application_map.py and
references/scoping_methodology.md §6). If no row in the whole input has
this field, the column is omitted entirely rather than left empty.

Optional "technical_blocker" boolean field per row (see
references/scoping_methodology.md §3) adds a "Technical Blocker" column,
flagging items that need engineering remediation regardless of Category
(a Simplification Item consistency error) rather than a business
decision — this is what tells sap-workshop-prep NOT to put the item in a
Fit-to-Standard discussion agenda. Same all-or-nothing column rule as
"applications": omitted entirely if no row sets it.
"""
import sys
import json
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

STATUS_FILL = {
    "Changed": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "Unchanged": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    "To Assess": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
    "Out of Scope": PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid"),
}
BLOCKER_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BLOCKER_FONT = Font(color="9C0006", bold=True)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)


def sheet_safe_name(name, used):
    base = name[:31]
    candidate = base
    i = 2
    while candidate in used:
        suffix = f" ({i})"
        candidate = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def build(input_path, output_path):
    with open(input_path) as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    summary_ws["A1"] = data.get("title", "Business Process Scope Summary")
    summary_ws["A1"].font = TITLE_FONT

    headers = ["Business Area", "# Processes", "Changed", "Unchanged", "To Assess", "Out of Scope"]
    for i, h in enumerate(headers, start=1):
        c = summary_ws.cell(row=3, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    totals = Counter()
    used_names = {"Summary"}
    row_idx = 4

    for area, rows in data["areas"].items():
        status_counts = Counter(r.get("status", "") for r in rows)
        summary_ws.cell(row=row_idx, column=1, value=area)
        summary_ws.cell(row=row_idx, column=2, value=len(rows))
        summary_ws.cell(row=row_idx, column=3, value=status_counts.get("Changed", 0))
        summary_ws.cell(row=row_idx, column=4, value=status_counts.get("Unchanged", 0))
        summary_ws.cell(row=row_idx, column=5, value=status_counts.get("To Assess", 0))
        summary_ws.cell(row=row_idx, column=6, value=status_counts.get("Out of Scope", 0))
        row_idx += 1
        totals["processes"] += len(rows)
        for k in ("Changed", "Unchanged", "To Assess", "Out of Scope"):
            totals[k] += status_counts.get(k, 0)

    total_row = row_idx
    summary_ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    summary_ws.cell(row=total_row, column=2, value=totals["processes"]).font = Font(bold=True)
    summary_ws.cell(row=total_row, column=3, value=totals["Changed"]).font = Font(bold=True)
    summary_ws.cell(row=total_row, column=4, value=totals["Unchanged"]).font = Font(bold=True)
    summary_ws.cell(row=total_row, column=5, value=totals["To Assess"]).font = Font(bold=True)
    summary_ws.cell(row=total_row, column=6, value=totals["Out of Scope"]).font = Font(bold=True)

    note_row = total_row + 2
    for note in data.get("notes", []):
        summary_ws.cell(row=note_row, column=1, value=note).font = Font(italic=True, size=9)
        summary_ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
        note_row += 1

    for col in range(1, 7):
        summary_ws.column_dimensions[get_column_letter(col)].width = 22 if col == 1 else 14

    has_apps = any(row.get("applications") for rows in data["areas"].values() for row in rows)
    has_blocker = any(row.get("technical_blocker") for rows in data["areas"].values() for row in rows)

    for area, rows in data["areas"].items():
        ws = wb.create_sheet(sheet_safe_name(area, used_names))
        ws["A1"] = f"Business Process Scope: {area}"
        ws["A1"].font = TITLE_FONT

        col_headers = ["Stream", "Process", "Status", "Complexity"]
        if has_blocker:
            col_headers.append("Technical Blocker")
        col_headers.append("Driving Simplification Item / Note")
        if has_apps:
            col_headers.append("Supporting Application(s)")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(col_headers))
        for i, h in enumerate(col_headers, start=1):
            c = ws.cell(row=3, column=i, value=h)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT

        r = 4
        for row in rows:
            col = 1
            ws.cell(row=r, column=col, value=row.get("stream", "")); col += 1
            ws.cell(row=r, column=col, value=row.get("process", "")); col += 1
            status = row.get("status", "")
            status_cell = ws.cell(row=r, column=col, value=status)
            if status in STATUS_FILL:
                status_cell.fill = STATUS_FILL[status]
            col += 1
            ws.cell(row=r, column=col, value=row.get("complexity", "")); col += 1
            if has_blocker:
                is_blocker = bool(row.get("technical_blocker"))
                blocker_cell = ws.cell(row=r, column=col, value="Yes — pre-workshop fix" if is_blocker else "")
                if is_blocker:
                    blocker_cell.fill = BLOCKER_FILL
                    blocker_cell.font = BLOCKER_FONT
                col += 1
            note = row.get("driving_item", "")
            if row.get("note"):
                note = f"{note} — {row['note']}" if note else row["note"]
            ws.cell(row=r, column=col, value=note); col += 1
            if has_apps:
                apps = row.get("applications", "")
                if isinstance(apps, list):
                    apps = ", ".join(apps)
                ws.cell(row=r, column=col, value=apps); col += 1
            r += 1

        base_widths = {"Stream": 28, "Process": 40, "Status": 14, "Complexity": 14,
                       "Technical Blocker": 20, "Driving Simplification Item / Note": 55,
                       "Supporting Application(s)": 30}
        for i, h in enumerate(col_headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = base_widths.get(h, 20)
        ws.freeze_panes = "A4"

    wb.save(output_path)
    print(f"Wrote workbook with {len(data['areas'])} area sheets + Summary to {output_path}")
    print(f"Totals: {dict(totals)}")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])


if __name__ == "__main__":
    main()
